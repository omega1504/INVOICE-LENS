
import re
import ollama
import pandas as pd
import streamlit as st

from pypdf import PdfReader
from docx import Document


# =========================================================
# PAGE SETTINGS
# =========================================================

st.set_page_config(
    page_title="InvoiceLens",
    page_icon="🐬",
    layout="wide"
)

st.title("InvoiceLens")
st.caption("Your invoices, understood.")
st.divider()


# =========================================================
# COMMON INVOICE SCHEMA
# =========================================================

INVOICE_COLUMNS = [
    "Invoice No",
    "Vendor Name",
    "Customer Name",
    "Invoice Date",
    "Due Date",
    "Amount",
    "GST / Tax",
    "Payment Terms",
    "Status",
    "Page"
]


# =========================================================
# READ PDF
# =========================================================

def read_pdf(file):
    reader = PdfReader(file)
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        pages.append({
            "page": page_number,
            "text": text
        })

    return pages


# =========================================================
# READ WORD
# =========================================================

def read_word(file):
    document = Document(file)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(values):
                parts.append(" | ".join(values))

    return "\n".join(parts)


# =========================================================
# READ EXCEL
# =========================================================

def read_excel(file):
    return pd.read_excel(file)


# =========================================================
# SAFE MISSING VALUE CHECK
# =========================================================

def is_missing(value):
    if value is None:
        return True

    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


# =========================================================
# CLEAN AMOUNT
# =========================================================

def clean_amount(value):
    if is_missing(value):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    value = str(value).strip()

    if not value:
        return None

    value = re.sub(
        r"(?i)\b(?:INR|Rs\.?)\s*",
        "",
        value
    )

    value = (
        value
        .replace("₹", "")
        .replace("$", "")
        .replace("€", "")
        .replace("£", "")
        .replace(",", "")
        .replace("■", "")
        .strip()
    )

    match = re.search(
        r"-?\d+(?:\.\d+)?",
        value
    )

    if not match:
        return None

    try:
        return float(match.group())
    except (TypeError, ValueError):
        return None


# =========================================================
# CLEAN DATE
# =========================================================

def clean_date(value):
    if is_missing(value):
        return pd.NaT

    return pd.to_datetime(
        value,
        errors="coerce",
        dayfirst=True
    )


# =========================================================
# FIND INVOICE NUMBER
# =========================================================

def find_invoice_number(text):
    patterns = [
        r"\bInvoice\s+No\.?\s*[:#\-]?\s*([A-Za-z0-9][A-Za-z0-9./_-]*)",
        r"\bInvoice\s+Number\s*[:#\-]?\s*([A-Za-z0-9][A-Za-z0-9./_-]*)",
        r"\bInvoice\s+#\s*[:\-]?\s*([A-Za-z0-9][A-Za-z0-9./_-]*)",
        r"\bInvoice\s+ID\s*[:#\-]?\s*([A-Za-z0-9][A-Za-z0-9./_-]*)",
        r"\bTax\s+Invoice\s+(?!Date\b)([A-Za-z0-9][A-Za-z0-9./_-]*)"
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:
            value = match.group(1).strip()

            if value.lower() not in {
                "date", "number", "no"
            }:
                return value

    fallback_patterns = [
        r"\b[A-Z]{2,10}/\d{4}[-/]\d+\b",
        r"\b[A-Z]{2,10}-\d{3,}\b",
        r"\b[A-Z]{1,10}/\d{2,4}/\d{2,}\b",
        r"\b[A-Z]{2,10}_\d{2,4}_\d{2,}\b"
    ]

    for pattern in fallback_patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:
            return match.group(0)

    return None


# =========================================================
# FIND DATE
# =========================================================

def find_date(text, labels):
    for label in labels:
        pattern = (
            rf"\b{re.escape(label)}\b"
            r"\s*[:\-]?\s*"
            r"(\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4}"
            r"|\d{4}[-/.]\d{1,2}[-/.]\d{1,2})"
        )

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:
            return clean_date(match.group(1))

    return pd.NaT


# =========================================================
# FIND LABELED VALUE
# =========================================================

def find_value(text, labels, stop_labels=None):
    stop_labels = stop_labels or []

    for label in labels:
        if stop_labels:
            stops = "|".join(
                re.escape(item)
                for item in stop_labels
            )

            pattern = (
                rf"\b{re.escape(label)}\b"
                rf"\s*[:\-]?\s*(.*?)"
                rf"(?=\s+(?:{stops})\b|$)"
            )
        else:
            pattern = (
                rf"\b{re.escape(label)}\b"
                rf"\s*[:\-]?\s*([^\n|]+)"
            )

        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:
            value = re.sub(
                r"\s+",
                " ",
                match.group(1)
            ).strip()

            if value:
                return value

    return None


# =========================================================
# FIND CUSTOMER NAME
# =========================================================

def find_customer(text):
    patterns = [
        r"\bCustomer\s+Name\s*[:\-]\s*([^\n|]+)",
        r"\bClient\s+Name\s*[:\-]\s*([^\n|]+)",
        r"\bBuyer\s+Name\s*[:\-]\s*([^\n|]+)",
        r"\bCustomer\s*[:\-]\s*([^\n|]+)",
        r"\bClient\s*[:\-]\s*([^\n|]+)",
        r"\bBuyer\s*[:\-]\s*([^\n|]+)"
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:
            value = match.group(1).strip()

            value = re.split(
                r"\s+(?:Address|GSTIN|Phone|Email)\s*:",
                value,
                flags=re.IGNORECASE
            )[0].strip()

            if value and value.lower() not in {
                "unknown",
                "customer",
                "customer name"
            }:
                return value

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    for i, line in enumerate(lines):
        if re.fullmatch(
            r"Bill\s+To\s*:?",
            line,
            flags=re.IGNORECASE
        ):
            if i + 1 < len(lines):
                candidate = lines[i + 1]

                # Handle "Bill To" followed by
                # "Customer Name: ABC Traders".
                candidate = re.sub(
                    r"^(?:Customer|Client|Buyer)\s+Name\s*:\s*",
                    "",
                    candidate,
                    flags=re.IGNORECASE
                ).strip()

                if candidate and not re.match(
                    r"^(Address|GSTIN|Phone|Email)\s*:",
                    candidate,
                    flags=re.IGNORECASE
                ):
                    return candidate

    return "Unknown"


# =========================================================
# FIND VENDOR NAME
# =========================================================

def find_vendor(text):
    patterns = [
        r"\bVendor\s+Name\s*[:\-]\s*([^\n|]+)",
        r"\bSupplier\s+Name\s*[:\-]\s*([^\n|]+)",
        r"\bSeller\s+Name\s*[:\-]\s*([^\n|]+)",
        r"\bVendor\s*[:\-]\s*([^\n|]+)",
        r"\bSupplier\s*[:\-]\s*([^\n|]+)",
        r"\bSeller\s*[:\-]\s*([^\n|]+)"
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:
            value = match.group(1).strip()

            value = re.split(
                r"\s+(?:Address|GSTIN|Phone|Email)\s*:",
                value,
                flags=re.IGNORECASE
            )[0].strip()

            if value:
                return value

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    # Example: AX INVOICE NANDI OFFICE SUPPLIES
    for line in lines[:12]:
        match = re.match(
            r"^\s*(?:AX\s+)?(?:TAX\s+)?INVOICE\s+(.+?)\s*$",
            line,
            flags=re.IGNORECASE
        )

        if match:
            value = match.group(1).strip(" :-|")

            if value:
                return value

    # Check for a company heading immediately before
    # invoice metadata.
    for i, line in enumerate(lines):
        if re.search(
            r"\bInvoice\s+No\b",
            line,
            flags=re.IGNORECASE
        ):
            for candidate in reversed(
                lines[max(0, i - 5):i]
            ):
                if not re.search(
                    r"\b(invoice|date|due date|payment terms|gstin)\b",
                    candidate,
                    flags=re.IGNORECASE
                ):
                    if not re.fullmatch(
                        r"[\d./_-]+",
                        candidate
                    ):
                        return candidate

    return "Unknown"


# =========================================================
# FIND TOTAL AMOUNT
# =========================================================

def find_amount(text):
    """Extract the final invoice total from common invoice labels."""
    text = re.sub(r"\s+", " ", text or "").strip()

    currency = r"(?:INR|Rs\.?|₹|\$|€|£)?\s*"
    number = r"([\d,]+(?:\.\d{1,2})?)"
    patterns = [
        rf"\bGrand\s+Total(?:\s+Amount)?\s*[:\-]?\s*{currency}{number}",
        rf"\bInvoice\s+Total(?:\s+Amount)?\s*[:\-]?\s*{currency}{number}",
        rf"\bTotal\s+Amount\s*[:\-]?\s*{currency}{number}",
        rf"\bAmount\s+Due\s*[:\-]?\s*{currency}{number}",
        rf"\bNet\s+Amount\s*[:\-]?\s*{currency}{number}",
        rf"\bTotal\s*:\s*{currency}{number}",
    ]

    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return clean_amount(match.group(1))

    return None


# =========================================================
# FIND GST / TAX
# =========================================================

def find_gst(text):
    """Extract GST/tax values, including labels with rates in parentheses."""
    text = re.sub(r"\s+", " ", text or "").strip()

    currency = r"(?:INR|Rs\.?|₹|\$|€|£)?\s*"
    number = r"([\d,]+(?:\.\d{1,2})?)"
    optional_rate = r"(?:\s*\([^)]*\))?"
    patterns = [
        rf"\bGST\s*/\s*Tax\s+Amount{optional_rate}\s*[:\-]?\s*{currency}{number}",
        rf"\bGST\s+Amount{optional_rate}\s*[:\-]?\s*{currency}{number}",
        rf"\bTax\s+Amount{optional_rate}\s*[:\-]?\s*{currency}{number}",
        rf"\bGST\s*@\s*\d+(?:\.\d+)?\s*%\s*[:\-]\s*{currency}{number}",
        rf"\bGST\s*[:\-]\s*{currency}{number}",
        rf"\bTax\s*[:\-]\s*{currency}{number}",
    ]

    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return clean_amount(match.group(1))

    return None


# =========================================================
# FIND PAYMENT TERMS
# =========================================================

def find_payment_terms(text):
    return find_value(
        text,
        [
            "Payment Terms",
            "Payment Term",
            "Terms"
        ],
        [
            "S.No.",
            "S.No",
            "Description",
            "Quantity",
            "Amount",
            "Grand Total",
            "Invoice Date",
            "Due Date",
            "Payment Status"
        ]
    )


# =========================================================
# PAYMENT STATUS
# =========================================================

def normalize_status(status, due_date=None):
    if is_missing(status):
        value = ""
    else:
        value = str(status).strip().lower()

    # Explicit paid status
    if value in {"paid", "settled", "fully paid"}:
        return "Paid"

    # Explicit overdue status
    if value in {"overdue", "past due"}:
        return "Overdue"

    # Unpaid invoices
    if value in {
        "pending", "unpaid", "due",
        "partially paid", "part-paid"
    }:
        if (
            due_date is not None
            and not is_missing(due_date)
        ):
            due = clean_date(due_date)

            if not pd.isna(due):
                today = pd.Timestamp.today().normalize()

                if due.normalize() < today:
                    return "Overdue"

        return "Pending"

    # Use the due date only when no clear status exists.
    if due_date is not None and not is_missing(due_date):
        due = clean_date(due_date)

        if not pd.isna(due):
            today = pd.Timestamp.today().normalize()

            if due.normalize() < today:
                return "Overdue"

            return "Pending"

    return "Unknown"


def find_status(text, due_date=None):
    patterns = [
        r"\bPayment\s+Status\s*[:\-]?\s*(Paid|Pending|Unpaid|Overdue|Settled|Due|Partially\s+Paid)",
        r"\bInvoice\s+Status\s*[:\-]?\s*(Paid|Pending|Unpaid|Overdue|Settled|Due)",
        r"(?m)^\s*Status\s*[:\-]\s*(Paid|Pending|Unpaid|Overdue|Settled|Due)\s*$"
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        if match:
            return normalize_status(
                match.group(1),
                due_date
            )

    # Avoid interpreting "Thank you for paid service"
    # or unrelated text as payment confirmation.
    return normalize_status(
        "",
        due_date
    )


# =========================================================
# STANDARDIZE ALL INVOICES
# =========================================================

def standardize_invoice_df(df):
    if df is None or df.empty:
        return pd.DataFrame(columns=INVOICE_COLUMNS)

    df = df.copy()

    for column in INVOICE_COLUMNS:
        if column not in df.columns:
            df[column] = None

    for column in [
        "Vendor Name",
        "Customer Name",
        "Payment Terms"
    ]:
        df[column] = df[column].apply(
            lambda value: (
                "Unknown"
                if is_missing(value) or not str(value).strip()
                else str(value).strip()
            )
        )

    df["Status"] = [
        normalize_status(status, due_date)
        for status, due_date in zip(
            df["Status"],
            df["Due Date"]
        )
    ]

    df["Amount"] = df["Amount"].apply(clean_amount)
    df["GST / Tax"] = df["GST / Tax"].apply(clean_amount)

    df["Invoice Date"] = df["Invoice Date"].apply(clean_date)
    df["Due Date"] = df["Due Date"].apply(clean_date)

    return df[INVOICE_COLUMNS]


# =========================================================
# PARSE ONE PDF / WORD INVOICE
# =========================================================

def parse_invoice(text, page=None):
    invoice_no = find_invoice_number(text)

    if not invoice_no:
        return None

    invoice_date = find_date(
        text,
        ["Invoice Date", "Invoice Dt", "Bill Date"]
    )

    if pd.isna(invoice_date):
        invoice_date = find_date(text, ["Date"])

    due_date = find_date(
        text,
        ["Due Date", "Payment Due", "Due"]
    )

    vendor = find_vendor(text)
    customer = find_customer(text)

    amount = find_amount(text)
    gst = find_gst(text)

    payment_terms = find_payment_terms(text)
    status = find_status(text, due_date)

    return {
        "Invoice No": invoice_no,
        "Vendor Name": vendor,
        "Customer Name": customer,
        "Invoice Date": invoice_date,
        "Due Date": due_date,
        "Amount": amount,
        "GST / Tax": gst,
        "Payment Terms": payment_terms or "Unknown",
        "Status": status,
        "Page": page
    }


# =========================================================
# PROCESS PDF
# =========================================================

def process_pdf(file):
    pages = read_pdf(file)
    invoices = []

    extracted_text = []

    for page in pages:
        text = page["text"]

        extracted_text.append(
            f"--- PAGE {page['page']} ---\n{text}"
        )

        invoice = parse_invoice(
            text,
            page["page"]
        )

        if invoice:
            invoices.append(invoice)

    df = standardize_invoice_df(
        pd.DataFrame(invoices)
    )

    return df, "\n\n".join(extracted_text)


# =========================================================
# PROCESS WORD
# =========================================================

def process_word(file):
    text = read_word(file)

    extracted_text = text
    invoices = []

    # Locate invoice boundaries
    pattern = (
        r"\bInvoice\s+No\.?\s*[:#\-]?\s*"
        r"[A-Za-z0-9][A-Za-z0-9./_-]*"
        r"|\bInvoice\s+Number\s*[:#\-]?\s*"
        r"[A-Za-z0-9][A-Za-z0-9./_-]*"
    )

    matches = list(
        re.finditer(
            pattern,
            text,
            flags=re.IGNORECASE
        )
    )

    # Single invoice: keep the entire document, including
    # the vendor heading before Invoice No.
    if len(matches) <= 1:
        invoice = parse_invoice(text, 1)

        if invoice:
            invoices.append(invoice)

    else:
        # Preserve heading and vendor information for
        # the first invoice.
        for i, match in enumerate(matches):
            start = 0 if i == 0 else match.start()

            if i + 1 < len(matches):
                end = matches[i + 1].start()
            else:
                end = len(text)

            invoice_text = text[start:end]

            invoice = parse_invoice(
                invoice_text,
                i + 1
            )

            if invoice:
                invoices.append(invoice)

    df = standardize_invoice_df(
        pd.DataFrame(invoices)
    )

    return df, extracted_text


# =========================================================
# PROCESS EXCEL
# =========================================================

def process_excel(file):
    raw_df = read_excel(file)

    if raw_df.empty:
        return pd.DataFrame(columns=INVOICE_COLUMNS)

    df = raw_df.copy()

    # Normalize headers
    df.columns = [
        re.sub(
            r"_+",
            "_",
            re.sub(
                r"[^a-z0-9]+",
                "_",
                str(column).strip().lower()
            )
        ).strip("_")
        for column in df.columns
    ]

    def find_column(possible_names):
        for name in possible_names:
            if name in df.columns:
                return name

        return None

    invoice_column = find_column([
        "invoice_no",
        "invoice_number",
        "invoice",
        "invoice_id",
        "bill_no",
        "bill_number",
        "tax_invoice_no"
    ])

    vendor_column = find_column([
        "vendor",
        "vendor_name",
        "supplier",
        "supplier_name",
        "seller",
        "seller_name"
    ])

    customer_column = find_column([
        "customer",
        "customer_name",
        "client",
        "client_name",
        "buyer",
        "buyer_name",
        "bill_to"
    ])

    invoice_date_column = find_column([
        "invoice_date",
        "bill_date",
        "date"
    ])

    due_date_column = find_column([
        "due_date",
        "payment_due",
        "due"
    ])

    amount_column = find_column([
        "amount",
        "total",
        "total_amount",
        "invoice_amount",
        "grand_total",
        "net_amount",
        "total_invoice_amount"
    ])

    gst_column = find_column([
        "gst",
        "gst_amount",
        "tax",
        "tax_amount",
        "gst_tax"
    ])

    status_column = find_column([
        "status",
        "payment_status",
        "invoice_status",
        "payment_state"
    ])

    terms_column = find_column([
        "payment_terms",
        "payment_term",
        "terms"
    ])

    if invoice_column is None:
        # Do not mistake an arbitrary spreadsheet
        # for a list of invoices.
        return pd.DataFrame(columns=INVOICE_COLUMNS)

    result = pd.DataFrame(index=df.index)

    result["Invoice No"] = (
        df[invoice_column]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    # Vendor
    if vendor_column:
        result["Vendor Name"] = df[vendor_column]
    else:
        result["Vendor Name"] = "Unknown"

    # Customer
    if customer_column:
        result["Customer Name"] = df[customer_column]
    else:
        result["Customer Name"] = "Unknown"

    # Dates
    if invoice_date_column:
        result["Invoice Date"] = (
            df[invoice_date_column].apply(clean_date)
        )
    else:
        result["Invoice Date"] = pd.NaT

    if due_date_column:
        result["Due Date"] = (
            df[due_date_column].apply(clean_date)
        )
    else:
        result["Due Date"] = pd.NaT

    # Amount
    if amount_column:
        result["Amount"] = (
            df[amount_column].apply(clean_amount)
        )
    else:
        result["Amount"] = None

    # GST
    if gst_column:
        result["GST / Tax"] = (
            df[gst_column].apply(clean_amount)
        )
    else:
        result["GST / Tax"] = None

    # Payment status
    if status_column:
        result["Status"] = df[status_column]
    else:
        result["Status"] = "Unknown"

    # Payment terms
    if terms_column:
        result["Payment Terms"] = df[terms_column]
    else:
        result["Payment Terms"] = "Unknown"

    result["Page"] = None

    # Remove rows without invoice numbers
    result = result[
        result["Invoice No"].ne("")
        & result["Invoice No"].str.lower().ne("nan")
    ]

    return standardize_invoice_df(result)


# =========================================================
# AI QUESTIONS
# =========================================================

def ask_invoice_ai(question, df):
    ai_df = df.copy()

    for column in ["Invoice Date", "Due Date"]:
        if column in ai_df.columns:
            ai_df[column] = ai_df[column].astype(str)

    invoice_data = ai_df.to_dict(orient="records")

    system_prompt = """
You are InvoiceLens, an AI invoice analyst.

Answer business questions using ONLY the supplied invoice data.

Never invent company names, invoice numbers, amounts,
dates, or payment statuses.

Explain that status is Unknown when the data does not
contain enough information.

Distinguish vendor from customer:
- Vendor Name is the seller or supplier.
- Customer Name is the buyer or recipient.

Keep answers simple and business-friendly.
"""

    user_prompt = f"""
Invoice data:

{invoice_data}

User question:

{question}

Answer using only the invoice data.
"""

    try:
        response = ollama.chat(
            model="llama3.2",
            messages=[
                {
                    "role": "system",
                    "content": system_prompt
                },
                {
                    "role": "user",
                    "content": user_prompt
                }
            ]
        )

        return response["message"]["content"]

    except Exception as error:
        return (
            "InvoiceLens could not connect to the local AI.\n\n"
            f"Error: {error}"
        )


# =========================================================
# FILE UPLOAD
# =========================================================

uploaded_file = st.file_uploader(
    "Upload your invoice file",
    type=["pdf", "docx", "xlsx", "xls"]
)

if not uploaded_file:
    st.info("Upload a PDF, Word, or Excel invoice file.")
    st.stop()


# =========================================================
# PROCESS FILE
# =========================================================

file_type = uploaded_file.name.rsplit(".", 1)[-1].lower()

try:
    if file_type == "pdf":
        source = "PDF"

        df, extracted_text = process_pdf(uploaded_file)

        with st.expander("🔎 Show extracted PDF content"):
            st.text(extracted_text)

    elif file_type == "docx":
        source = "Word"

        df, extracted_text = process_word(uploaded_file)

        with st.expander("🔎 Show extracted Word content"):
            st.text(extracted_text)

    elif file_type in ["xlsx", "xls"]:
        source = "Excel"

        df = process_excel(uploaded_file)

    else:
        st.error("Unsupported file type.")
        st.stop()

except Exception as error:
    st.error(f"Could not read the {file_type.upper()} file: {error}")
    st.stop()


# =========================================================
# VALIDATE
# =========================================================

if df.empty:
    st.error(
        f"InvoiceLens read the {source} file, "
        "but could not identify any invoices."
    )

    st.info(
        "Check that the file contains invoice numbers "
        "or that the Excel sheet has an invoice number column."
    )

    st.stop()


df = standardize_invoice_df(df)

df = df[
    df["Invoice No"].astype(str).str.strip().ne("")
].copy()

if df.empty:
    st.error("No valid invoice numbers were found.")
    st.stop()


# =========================================================
# SUCCESS
# =========================================================

st.success(
    f"✓ {source} processed successfully — "
    f"{len(df)} invoice(s) detected."
)


# =========================================================
# DASHBOARD
# =========================================================

st.subheader("Overview")

total_invoices = len(df)

total_value = df["Amount"].fillna(0).sum()

paid_df = df[df["Status"] == "Paid"]
pending_df = df[df["Status"] == "Pending"]
overdue_df = df[df["Status"] == "Overdue"]
unknown_df = df[df["Status"] == "Unknown"]

paid_amount = paid_df["Amount"].fillna(0).sum()
pending_amount = pending_df["Amount"].fillna(0).sum()
overdue_amount = overdue_df["Amount"].fillna(0).sum()

outstanding_amount = pending_amount + overdue_amount

col1, col2, col3, col4 = st.columns(4)

col1.metric("Total Invoices", total_invoices)
col2.metric("Total Value", f"₹{total_value:,.2f}")
col3.metric("Outstanding", f"₹{outstanding_amount:,.2f}")
col4.metric("Overdue Invoices", len(overdue_df))


# =========================================================
# PAYMENT STATUS
# =========================================================

st.divider()
st.subheader("Payment Status")

c1, c2, c3, c4 = st.columns(4)

c1.metric(
    "Paid",
    len(paid_df),
    f"₹{paid_amount:,.2f}"
)

c2.metric(
    "Pending",
    len(pending_df),
    f"₹{pending_amount:,.2f}"
)

c3.metric(
    "Overdue",
    len(overdue_df),
    f"₹{overdue_amount:,.2f}"
)

c4.metric(
    "Unknown",
    len(unknown_df)
)

if len(overdue_df) > 0:
    st.warning(
        f"⚠️ {len(overdue_df)} invoice(s) are overdue."
    )

if len(unknown_df) > 0:
    st.info(
        f"ℹ️ {len(unknown_df)} invoice(s) have "
        "an undetermined payment status."
    )


# =========================================================
# INVOICE TABLE
# =========================================================

st.divider()
st.subheader("Invoices")

filter_status = st.selectbox(
    "Filter by payment status",
    ["All", "Paid", "Pending", "Overdue", "Unknown"]
)

if filter_status == "All":
    filtered_df = df.copy()
else:
    filtered_df = df[df["Status"] == filter_status].copy()

display_columns = [
    "Invoice No",
    "Vendor Name",
    "Customer Name",
    "Invoice Date",
    "Due Date",
    "Amount",
    "GST / Tax",
    "Payment Terms",
    "Status"
]

st.dataframe(
    filtered_df[display_columns],
    use_container_width=True,
    hide_index=True,
    column_config={
        "Invoice No": st.column_config.TextColumn(
            "Invoice Number"
        ),
        "Vendor Name": st.column_config.TextColumn(
            "Vendor"
        ),
        "Customer Name": st.column_config.TextColumn(
            "Customer"
        ),
        "Invoice Date": st.column_config.DateColumn(
            "Invoice Date",
            format="DD/MM/YYYY"
        ),
        "Due Date": st.column_config.DateColumn(
            "Due Date",
            format="DD/MM/YYYY"
        ),
        "Amount": st.column_config.NumberColumn(
            "Total Amount",
            format="₹%.2f"
        ),
        "GST / Tax": st.column_config.NumberColumn(
            "GST / Tax",
            format="₹%.2f"
        ),
        "Status": st.column_config.TextColumn(
            "Payment Status"
        )
    }
)


# =========================================================
# DOWNLOAD PROCESSED INVOICES
# =========================================================

st.divider()
st.subheader("Export Invoices")

csv_data = df[display_columns].to_csv(
    index=False
).encode("utf-8-sig")

st.download_button(
    "⬇️ Download Invoice Data (CSV)",
    data=csv_data,
    file_name="invoicelens_invoices.csv",
    mime="text/csv"
)


# =========================================================
# ASK INVOICELENS
# =========================================================

st.divider()
st.subheader("💬 Ask InvoiceLens")

st.caption(
    "Ask questions about your invoices in natural language."
)

question = st.text_input(
    "What would you like to know?",
    placeholder="Example: Who has not paid?"
)

if question:
    with st.spinner("Analyzing your invoices..."):
        answer = ask_invoice_ai(question, df)

    st.markdown(answer)